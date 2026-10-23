"""Read-only validation of the fixed English synthetic release."""
from pathlib import Path
from collections import Counter
import re
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
FILE = ROOT / 'workbook/Script_Store_Operations_Synthetic_Compatible.xlsx'

def main():
    with zipfile.ZipFile(FILE) as z:
        assert z.testzip() is None
        assert not any('externalLinks/' in n or 'pivotCache' in n for n in z.namelist())
        for n in z.namelist():
            if n.endswith('.xml'):
                assert not re.search(r'[\u4e00-\u9fff]', ' '.join(ET.fromstring(z.read(n)).itertext())), n
    for p in ROOT.rglob('*'):
        if p.is_file() and p.suffix in {'.md','.py','.txt'}:
            assert not re.search(r'[\u4e00-\u9fff]',p.read_text()),p
    w = openpyxl.load_workbook(FILE,data_only=True)
    raw = [r for r in list(w['Raw Data'].values)[1:] if r[0]]
    scripts = [r for r in list(w['Script List'].values)[1:] if r[0]]
    assert len(raw) == len({r[0] for r in raw}) == 416
    assert len(scripts) == len({r[0] for r in scripts}) == 160
    lookup = {r[0]:r for r in scripts}
    for r in raw:
        s = lookup[r[1]]
        assert r[4] == s[2] and r[10] == s[8]
        assert r[8] == r[6]*r[7] and r[9] == s[7] and r[11] == r[8]-r[9]
    assert {r[4] for r in raw} == {'Emotional','Deduction','Comedy','Mechanics'}
    assert {r[10] for r in raw} == {'City-limited','Exclusive','Boxed'}
    assert (sum(r[8] for r in raw),sum(r[9] for r in raw),sum(r[11] for r in raw)) == (280352,84912,195440)
    assert len(Counter(r[1] for r in raw)) == 136
    for s in scripts:
        rows = [r for r in raw if r[1] == s[0]]
        assert s[17] == len(rows) and s[18] == sum(r[11] for r in rows)
    assert all(w['Checks'].cell(r,4).value == 0 for r in range(2,11))
    for sheet in w:
        assert not any(c.data_type == 'e' for row in sheet for c in row),sheet.title
    assert len(w['Dashboard']._charts) == 3
    assert sum(len(s.tables) for s in w) == 2
    print('PASS: English content, unique keys, row calculations, aggregate reconciliation, cached errors, tables and charts.')
    print('416 sessions; 160 scripts; 136 opened; revenue $280,352; wages $84,912; contribution $195,440.')
    print('Cached-value checks do not replace native Excel recalculation or visual acceptance.')

if __name__ == '__main__':
    main()
